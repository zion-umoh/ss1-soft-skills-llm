"""Standardise the BFI-2/BESSI college-student workbook for Model 2.

The raw workbook stays unchanged. This script uses its supplied BFI-2 and
BESSI domain-scale columns, excludes rows missing any model input or target,
converts BFI-2 scores from their documented 1--5 response scale to 0--1, and
creates deterministic train/validation/test splits.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import random
import statistics
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


SPLITS = ("train", "validation", "test")
RAW_SCORE_MIN = 1.0
RAW_SCORE_MAX = 5.0
CASE_COLUMN = "Case"
BFI2_COLUMNS = {
    "BFI2_Extraversion": "big_five_extraversion",
    "BFI2_Agreeableness": "big_five_agreeableness",
    "BFI2_Conscientiousness": "big_five_conscientiousness",
    "BFI2_NegativeEmotionality": "big_five_negative_emotionality",
    "BFI2_OpenMindedness": "big_five_open_mindedness",
}
BFI2_FACET_COLUMNS = {
    "BFI2_E_Sociability": "big_five_extraversion_sociability",
    "BFI2_E_Assertiveness": "big_five_extraversion_assertiveness",
    "BFI2_E_EnergyLevel": "big_five_extraversion_energy_level",
    "BFI2_A_Compassion": "big_five_agreeableness_compassion",
    "BFI2_A_Respectfulness": "big_five_agreeableness_respectfulness",
    "BFI2_A_Trust": "big_five_agreeableness_trust",
    "BFI2_C_Organization": "big_five_conscientiousness_organization",
    "BFI2_C_Productiveness": "big_five_conscientiousness_productiveness",
    "BFI2_C_Responsibility": "big_five_conscientiousness_responsibility",
    "BFI2_N_Anxiety": "big_five_negative_emotionality_anxiety",
    "BFI2_N_Depression": "big_five_negative_emotionality_depression",
    "BFI2_N_EmotionalVolatility": "big_five_negative_emotionality_emotional_volatility",
    "BFI2_O_IntellectualCuriosity": "big_five_open_mindedness_intellectual_curiosity",
    "BFI2_O_AestheticSensitivity": "big_five_open_mindedness_aesthetic_sensitivity",
    "BFI2_O_CreativeImagination": "big_five_open_mindedness_creative_imagination",
}
BESSI_COLUMNS = {
    "BESSI_SelfManagementSkills": "bessi_self_management",
    "BESSI_SocialEngagementSkills": "bessi_social_engagement",
    "BESSI_CooperationSkills": "bessi_cooperation",
    "BESSI_EmotionalResilienceSkills": "bessi_emotional_resilience",
    "BESSI_InnovationSkills": "bessi_innovation",
}
FEATURE_COLUMNS = tuple(BFI2_COLUMNS.values())
TARGET_COLUMNS = tuple(BESSI_COLUMNS.values())
PROCESSED_COLUMNS = ("record_id", *FEATURE_COLUMNS, *TARGET_COLUMNS, "split")


class PreparationError(ValueError):
    """Raised when the workbook does not meet the Model 2 preparation contract."""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def numeric_score(value: object, column: str, source_row: int) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise PreparationError(f"Row {source_row}: {column} must be numeric, not boolean.")
    try:
        score = float(value)
    except (TypeError, ValueError) as error:
        raise PreparationError(f"Row {source_row}: {column} must be numeric, got {value!r}.") from error
    if not math.isfinite(score) or not RAW_SCORE_MIN <= score <= RAW_SCORE_MAX:
        raise PreparationError(
            f"Row {source_row}: {column}={score!r} is outside the expected 1--5 scale."
        )
    return score


def target_counts(total: int, ratios: dict[str, float]) -> dict[str, int]:
    raw = {split: total * ratios[split] for split in SPLITS}
    counts = {split: math.floor(raw[split]) for split in SPLITS}
    remainder = total - sum(counts.values())
    for split in sorted(SPLITS, key=lambda split: (raw[split] - counts[split], split), reverse=True)[:remainder]:
        counts[split] += 1
    return counts


def write_csv(path: Path, fieldnames: Iterable[str], rows: Iterable[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def read_records(
    input_path: Path, feature_mapping: dict[str, str] = BFI2_COLUMNS
) -> tuple[list[dict[str, object]], list[dict[str, object]], dict[str, object]]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    if "College student sample" not in workbook.sheetnames:
        raise PreparationError("Workbook must contain a 'College student sample' worksheet.")
    worksheet = workbook["College student sample"]
    rows = worksheet.iter_rows(values_only=True)
    header = list(next(rows, ()))
    expected_columns = (CASE_COLUMN, *feature_mapping, *BESSI_COLUMNS)
    missing_columns = [column for column in expected_columns if column not in header]
    if missing_columns:
        raise PreparationError(f"Workbook is missing required columns: {missing_columns}.")
    indices = {column: header.index(column) for column in expected_columns}

    records: list[dict[str, object]] = []
    exclusions: list[dict[str, object]] = []
    source_cases: list[object] = []
    for source_row, row in enumerate(rows, start=2):
        case_value = row[indices[CASE_COLUMN]]
        if case_value is None or case_value == "":
            exclusions.append({"source_row_number": source_row, "source_case": "", "exclusion_reason": "missing_case"})
            continue
        source_cases.append(case_value)
        raw_features = {
            canonical: numeric_score(row[indices[source]], source, source_row)
            for source, canonical in feature_mapping.items()
        }
        raw_targets = {
            canonical: numeric_score(row[indices[source]], source, source_row)
            for source, canonical in BESSI_COLUMNS.items()
        }
        missing = [column for column, value in {**raw_features, **raw_targets}.items() if value is None]
        if missing:
            exclusions.append(
                {
                    "source_row_number": source_row,
                    "source_case": str(case_value),
                    "exclusion_reason": "missing_model_score:" + ",".join(missing),
                }
            )
            continue
        record_id = "bfi2_bessi_" + hashlib.sha256(
            f"bfi2-bessi:{source_row}".encode("utf-8")
        ).hexdigest()[:16]
        records.append(
            {
                "source_row_number": source_row,
                "source_case": str(case_value),
                "record_id": record_id,
                **{column: (value - RAW_SCORE_MIN) / (RAW_SCORE_MAX - RAW_SCORE_MIN) for column, value in raw_features.items()},
                **raw_targets,
            }
        )

    if not records:
        raise PreparationError("No rows have complete BFI-2 and BESSI domain scores.")
    case_counts = Counter(source_cases)
    linkage = {
        "raw_rows": worksheet.max_row - 1,
        "unique_case_values": len(case_counts),
        "reused_case_values": sum(count > 1 for count in case_counts.values()),
        "rows_with_reused_case_values": sum(count for count in case_counts.values() if count > 1),
        "case_identifier_is_unique": all(count == 1 for count in case_counts.values()),
    }
    return records, exclusions, {"header": [str(value) if value is not None else "" for value in header], "linkage": linkage}


def assign_splits(records: list[dict[str, object]], ratios: dict[str, float], seed: int) -> dict[str, list[dict[str, object]]]:
    ordered_records = list(records)
    random.Random(seed).shuffle(ordered_records)
    counts = target_counts(len(ordered_records), ratios)
    assignments: dict[str, list[dict[str, object]]] = {}
    start = 0
    for split in SPLITS:
        assigned = ordered_records[start : start + counts[split]]
        for record in assigned:
            record["split"] = split
        assignments[split] = assigned
        start += counts[split]
    return assignments


def distribution(records: list[dict[str, object]], columns: Iterable[str]) -> dict[str, dict[str, float]]:
    result: dict[str, dict[str, float]] = {}
    for column in columns:
        values = [float(record[column]) for record in records]
        result[column] = {
            "min": round(min(values), 4),
            "mean": round(statistics.fmean(values), 4),
            "median": round(statistics.median(values), 4),
            "max": round(max(values), 4),
        }
    return result


def make_summary(
    input_path: Path,
    records: list[dict[str, object]],
    exclusions: list[dict[str, object]],
    source_metadata: dict[str, object],
    assignments: dict[str, list[dict[str, object]]],
    seed: int,
    ratios: dict[str, float],
    feature_columns: tuple[str, ...],
    feature_level: str,
) -> dict[str, object]:
    return {
        "dataset_id": "bfi2_bessi_college_student_sample",
        "generated_at_utc": datetime.now(UTC).isoformat().replace("+00:00", "Z"),
        "source_file": input_path.as_posix(),
        "source_sha256": sha256(input_path),
        "source_schema": source_metadata["header"],
        "row_counts": {
            "raw_rows": source_metadata["linkage"]["raw_rows"],
            "eligible_rows": len(records),
            "excluded_rows": len(exclusions),
        },
        "participant_linkage": source_metadata["linkage"],
        "score_handling": {
            "source": f"supplied BFI-2 {feature_level}-scale columns and BESSI domain-scale columns",
            "source_scale": "1--5",
            "feature_conversion": "BFI-2 features are transformed to unit interval with (score - 1) / 4.",
            "target_scale": "BESSI targets remain on the original 1--5 scale for interpretability.",
            "excluded_for_missing_scores": True,
        },
        "preparation_policy": {
            "random_seed": seed,
            "split_ratios": ratios,
            "split_unit": "row",
            "split_grouping": "No trustworthy unique participant identifier is available: the Case field is reused for distinct measurement rows, so the split is row-level.",
        },
        "features": list(feature_columns),
        "feature_level": feature_level,
        "targets": list(TARGET_COLUMNS),
        "distributions": {
            "features_unit_interval": distribution(records, feature_columns),
            "targets_original_1_to_5": distribution(records, TARGET_COLUMNS),
        },
        "splits": {split: {"row_count": len(assignments[split])} for split in SPLITS},
    }


def render_report(summary: dict[str, object]) -> str:
    linkage = summary["participant_linkage"]
    policy = summary["preparation_policy"]
    rows = summary["row_counts"]
    lines = [
        "# BFI-2 and BESSI Dataset Quality Report",
        "",
        "## Standardisation decision",
        "",
        f"- Raw source: `{summary['source_file']}`",
        f"- SHA-256: `{summary['source_sha256']}`",
        f"- BFI-2 inputs: supplied {summary['feature_level']} scores, mapped from 1--5 to 0--1 using `(score - 1) / 4`.",
        "- BESSI targets: supplied domain scores retained on their original 1--5 scale.",
        f"- Random seed: `{policy['random_seed']}`; split ratios: `{policy['split_ratios']}`.",
        "",
        "## Data quality",
        "",
        f"- Raw rows: {rows['raw_rows']}",
        f"- Eligible model rows: {rows['eligible_rows']}",
        f"- Excluded rows with missing model scores: {rows['excluded_rows']}",
        f"- Unique values in `Case`: {linkage['unique_case_values']}",
        f"- Reused `Case` values: {linkage['reused_case_values']} across {linkage['rows_with_reused_case_values']} rows.",
        "",
        "## Linkage limitation",
        "",
        "`Case` is not a unique participant ID: repeated values have different measurement values. It therefore cannot be used to prove participant linkage or form participant-level splits. The prepared splits are row-level and this limitation must be reported in Model 2 evaluation. Do not claim participant-independent test performance unless a reliable unique participant identifier is supplied.",
        "",
        "## Standardised Model 2 schema",
        "",
        "| Role | Columns | Scale |",
        "| --- | --- | --- |",
        f"| Inputs | {', '.join(summary['features'])} | 0--1 |",
        f"| Targets | {', '.join(summary['targets'])} | 1--5 |",
        "",
        "## Score distributions",
        "",
        "| Variable | Min | Median | Mean | Max |",
        "| --- | ---: | ---: | ---: | ---: |",
    ]
    for label, statistics_by_column in summary["distributions"].items():
        for column, values in statistics_by_column.items():
            lines.append(
                f"| {column} ({label}) | {values['min']:.4f} | {values['median']:.4f} | {values['mean']:.4f} | {values['max']:.4f} |"
            )
    lines.extend(["", "## Fixed splits", "", "| Split | Rows |", "| --- | ---: |"])
    for split in SPLITS:
        lines.append(f"| {split} | {summary['splits'][split]['row_count']} |")
    lines.extend(
        [
            "",
            "## Modelling guardrails",
            "",
            "- Use only the processed split files for Model 2 modelling.",
            "- Do not use `source_case`, `source_row_number`, age, gender, ethnicity, race or other questionnaire fields as features.",
            "- Fit any additional scaler or imputer on the training split only.",
            "- Do not use the held-out test split for model selection or hyperparameter tuning.",
            "",
        ]
    )
    return "\n".join(lines)


def prepare(
    input_path: Path,
    interim_dir: Path,
    processed_dir: Path,
    report_path: Path,
    summary_path: Path,
    seed: int = 42,
    train_ratio: float = 0.70,
    validation_ratio: float = 0.15,
    test_ratio: float = 0.15,
    feature_mapping: dict[str, str] = BFI2_COLUMNS,
    feature_level: str = "domain",
) -> dict[str, object]:
    ratios = {"train": train_ratio, "validation": validation_ratio, "test": test_ratio}
    if any(value <= 0 for value in ratios.values()) or not math.isclose(sum(ratios.values()), 1.0):
        raise PreparationError("Split ratios must be positive and sum to 1.0.")
    feature_columns = tuple(feature_mapping.values())
    processed_columns = ("record_id", *feature_columns, *TARGET_COLUMNS, "split")
    records, exclusions, source_metadata = read_records(input_path, feature_mapping)
    assignments = assign_splits(records, ratios, seed)
    summary = make_summary(
        input_path,
        records,
        exclusions,
        source_metadata,
        assignments,
        seed,
        ratios,
        feature_columns,
        feature_level,
    )

    audit_columns = ("source_row_number", "source_case", *processed_columns)
    write_csv(interim_dir / "bfi2_bessi_cleaned_audit.csv", audit_columns, records)
    write_csv(
        interim_dir / "bfi2_bessi_exclusions.csv",
        ("source_row_number", "source_case", "exclusion_reason"),
        exclusions,
    )
    for split in SPLITS:
        write_csv(
            processed_dir / f"bfi2_bessi_{split}.csv",
            processed_columns,
            ({column: record[column] for column in processed_columns} for record in assignments[split]),
        )
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(render_report(summary), encoding="utf-8")
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input", type=Path, default=Path("data/raw/bfi2-bessi/College student sample.xlsx")
    )
    parser.add_argument("--interim-dir", type=Path, default=Path("data/interim/bfi2-bessi"))
    parser.add_argument("--processed-dir", type=Path, default=Path("data/processed/model2"))
    parser.add_argument(
        "--report", type=Path, default=Path("reports/data-quality/bfi2-bessi-data-quality-report.md")
    )
    parser.add_argument(
        "--summary", type=Path, default=Path("data/metadata/bfi2-bessi-preparation-summary.json")
    )
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()
    summary = prepare(
        input_path=args.input,
        interim_dir=args.interim_dir,
        processed_dir=args.processed_dir,
        report_path=args.report,
        summary_path=args.summary,
        seed=args.seed,
    )
    print(
        f"Prepared {summary['row_counts']['eligible_rows']} Model 2 rows across "
        f"{', '.join(f'{split}={summary['splits'][split]['row_count']}' for split in SPLITS)}."
    )


if __name__ == "__main__":
    main()
