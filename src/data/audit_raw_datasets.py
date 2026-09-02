"""Create a repeatable inventory of raw tabular datasets.

The audit never modifies files beneath ``data/raw``. It records each CSV, TSV
or XLSX file's path, byte size, modification time, SHA-256 checksum, schema,
and row count.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sys
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def delimited_summary(path: Path) -> tuple[list[str], int]:
    csv.field_size_limit(sys.maxsize)
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream, delimiter="\t" if path.suffix == ".tsv" else ",")
        if path.suffix == ".tsv" and "official_splits" in path.parts:
            return ["text", "emotion_ids", "comment_id"], sum(1 for _ in reader)
        header = next(reader, None)
        if header is None:
            return [], 0
        return header, sum(1 for _ in reader)


def workbook_summary(path: Path) -> tuple[list[str], int, dict[str, object]]:
    """Return the first-sheet schema and all-sheet dimensions without modifying XLSX."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets: dict[str, object] = {}
    first_header: list[str] = []
    first_row_count = 0
    for sheet_index, worksheet in enumerate(workbook.worksheets):
        header = ["" if value is None else str(value) for value in next(worksheet.values, ())]
        row_count = max(worksheet.max_row - (1 if header else 0), 0)
        sheets[worksheet.title] = {
            "row_count": row_count,
            "column_count": worksheet.max_column,
            "columns": header,
        }
        if sheet_index == 0:
            first_header, first_row_count = header, row_count
    return first_header, first_row_count, sheets


def audit(data_root: Path) -> list[dict[str, object]]:
    entries: list[dict[str, object]] = []
    paths = sorted(
        path
        for path in data_root.rglob("*")
        if path.is_file() and path.suffix.lower() in {".csv", ".tsv", ".xlsx"}
    )
    for path in paths:
        if path.suffix.lower() == ".xlsx":
            header, row_count, sheets = workbook_summary(path)
        else:
            header, row_count = delimited_summary(path)
            sheets = None
        stat = path.stat()
        entries.append(
            {
                "relative_path": path.as_posix(),
                "format": path.suffix.removeprefix("."),
                "file_size_bytes": stat.st_size,
                "modified_at_utc": datetime.fromtimestamp(stat.st_mtime, UTC)
                .isoformat()
                .replace("+00:00", "Z"),
                "sha256": sha256(path),
                "row_count": row_count,
                "column_count": len(header),
                "columns": header,
                **({"sheets": sheets} if sheets is not None else {}),
            }
        )
    return entries


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-root", type=Path, default=Path("data/raw"))
    parser.add_argument(
        "--output", type=Path, default=Path("data/metadata/raw-data-inventory.json")
    )
    args = parser.parse_args()

    if not args.data_root.is_dir():
        raise SystemExit(f"Raw-data directory not found: {args.data_root}")

    output = {
        "generated_at_utc": datetime.now(UTC).isoformat().replace("+00:00", "Z"),
        "data_root": args.data_root.as_posix(),
        "files": audit(args.data_root),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(output['files'])} entries to {args.output}")


if __name__ == "__main__":
    main()
